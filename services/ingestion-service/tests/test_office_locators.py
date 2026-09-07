"""Real OOXML fixtures through parsing, chunking and the ingestion index boundary."""

import asyncio
import hashlib
import io
import re
import zipfile
from dataclasses import replace

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.config import load_settings
from app.object_storage import SourceObject
from app.registry_client import RegistryClient
from chunkers.logical import LogicalStructureChunker
from parsers.base import ParserError
from parsers.office_limits import OfficeLimits
from parsers.pptx import PptxParser
from parsers.router import ParserRouter
from parsers.xlsx import XlsxParser
from tests.conftest import actor_proof_headers, make_client, web_transport_headers


def source(name, content):
    return SourceObject(uri=f"s3://test/{name}", filename=name, mime_type="application/octet-stream", content=content,
                        sha256="sha256:" + hashlib.sha256(content).hexdigest())


def workbook_source(*, many=False, bad_dimension=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rozpočet"
    for number, cells in [(5, ["ID", "Minimum", "Proposed", "Guaranteed"]),
                          (60, ["ROW_60", None, "32 GB", None]), (115, ["ROW_115", "End", None, "Verified"])]:
        for column, value in enumerate(cells, start=1):
            sheet.cell(number, column, value)
    if many:
        for number in range(116, 198):
            sheet.cell(number, 1, f"ROW_{number}")
            sheet.cell(number, 3, "Meaningful source evidence that must stay on its own source row.")
    workbook.create_sheet("Other").append(["OTHER_SHEET", "Evidence"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    if bad_dimension:
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(output, "w") as patched:
            for entry in original.infolist():
                body = original.read(entry)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    body = re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', body)
                patched.writestr(entry, body)
        content = output.getvalue()
    return source("budget.xlsx", content)


def presentation_source(*, long=False):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Sizing"
    for table_index in range(2):
        rows = 16 if long else 3
        table = slide.shapes.add_table(rows, 4, Inches(1), Inches(1+table_index), Inches(5), Inches(1)).table
        for cell, label in zip(table.rows[0].cells, ["Service", "Minimum", "Proposed", "Guaranteed"]):
            cell.text = label
        for number in range(1, rows):
            table.cell(number, 0).text = f"T{table_index}_ROW_{number+1}"
            table.cell(number, 2).text = "32 GB " + ("proposal evidence " * 4 if long else "")
    buffer = io.BytesIO()
    presentation.save(buffer)
    return source("sizing.pptx", buffer.getvalue())


def chunks(parsed, original, *, small=False):
    settings = load_settings({"AKL_ENV": "test", "AKL_AUTH_MODE": "disabled", "AKL_INGESTION_REGISTRY_CLIENT_MODE": "mock"})
    settings = replace(settings, chunk_target_chars=260 if small else 10000,
                       max_chunk_chars=320 if small else 12000, chunk_overlap_chars=20)
    metadata = asyncio.run(RegistryClient(settings).get_document_metadata("doc_office", "ver_office"))
    return LogicalStructureChunker(settings).chunk(
        parsed, document_metadata=metadata, source=original, extraction_profile="default",
        parser_profile="default", chunking_strategy="logical",
    ).chunks


@pytest.mark.parametrize("bad_dimension", [False, True])
def test_xlsx_sparse_late_rows_keep_physical_row_numbers_without_invented_pages(bad_dimension):
    original = workbook_source(bad_dimension=bad_dimension)
    parsed = XlsxParser().parse(original, parser_profile="default")
    assert parsed.pages_processed == 0
    assert parsed.metadata["sheets_processed"] == 2
    first = parsed.blocks[0]
    assert first.metadata["source_locator"]["row_numbers"] == [5, 60, 115]
    assert "ROW_60 |  | 32 GB | " in first.text
    assert "ROW_115" in first.text
    result = chunks(parsed, original)
    assert len(result) == 2  # Different worksheets never share the first locator.
    assert all(item.page_number is None for item in result)
    assert result[0].section_path == ["Rozpočet", "A:D · Řádky 5, 60, 115"]
    assert result[0].metadata["source_locator"]["sheet_name"] == "Rozpočet"
    assert "OTHER_SHEET" not in result[0].text


@pytest.mark.parametrize("kind", ["xlsx", "pptx"])
def test_split_tables_preserve_exact_rows_headers_columns_and_table_identity(kind):
    original = workbook_source(many=True) if kind == "xlsx" else presentation_source(long=True)
    parsed = (XlsxParser() if kind == "xlsx" else PptxParser()).parse(original, parser_profile="default")
    result = chunks(parsed, original, small=True)
    assert len(result) > 5
    all_rows = set()
    table_ids = set()
    for item in result:
        locator = item.metadata["source_locator"]
        assert item.page_number is None
        if item.metadata["block_type"] != "table":
            continue
        numbers = locator["row_numbers"]
        assert len(item.text.splitlines()) == len(numbers)
        assert len(item.text) <= 320
        if kind == "xlsx" and locator["sheet_name"] == "Rozpočet":
            assert item.text.splitlines()[0].startswith("ID | Minimum")
            assert numbers[0] == 5
            extracted = [int(value) for value in re.findall(r"ROW_(\d+)", item.text)]
            assert extracted == numbers[1:]
            all_rows.update(extracted)
        elif kind == "pptx":
            assert item.text.splitlines()[0] == "Service | Minimum | Proposed | Guaranteed"
            assert numbers[0] == 1
            extracted = [int(value) for value in re.findall(r"_ROW_(\d+)", item.text)]
            assert extracted == numbers[1:]
            assert all(" |  | 32 GB" in line and line.rstrip().endswith("|") for line in item.text.splitlines()[1:])
            table_ids.add(locator["table_id"])
    if kind == "xlsx":
        assert all_rows == {60, 115, *range(116, 198)}
    else:
        assert len(table_ids) == 2


@pytest.mark.parametrize("limit", [
    {"rows": 10}, {"columns": 2}, {"cells": 2}, {"text_chars": 10},
    {"zip_entries": 1}, {"zip_expanded_bytes": 16}, {"zip_entry_bytes": 16}, {"seconds": 0},
])
def test_workbook_processing_limits_reject_the_entire_result(limit):
    with pytest.raises(ParserError) as error:
        XlsxParser(replace(OfficeLimits(), **limit)).parse(workbook_source(), parser_profile="default")
    assert error.value.code == "XLSX_PROCESSING_LIMIT"


@pytest.mark.parametrize("limit", [{"slides": 0}, {"columns": 2}, {"cells": 2}, {"text_chars": 10}, {"rows": 1}])
def test_presentation_processing_limits_reject_the_entire_result(limit):
    with pytest.raises(ParserError) as error:
        PptxParser(replace(OfficeLimits(), **limit)).parse(presentation_source(), parser_profile="default")
    assert error.value.code == "PPTX_PROCESSING_LIMIT"


@pytest.mark.parametrize("kind", ["xlsx", "pptx"])
def test_oversized_table_row_is_rejected_instead_of_split_between_columns(kind):
    original = workbook_source() if kind == "xlsx" else presentation_source(long=True)
    parsed = (XlsxParser() if kind == "xlsx" else PptxParser()).parse(original, parser_profile="default")
    block = next(block for block in parsed.blocks if block.block_type == "table")
    oversized = replace(block, text=block.text.splitlines()[0] + "\n" + "x" * 1000,
                        metadata={**block.metadata, "source_locator": {**block.metadata["source_locator"], "row_numbers": [1, 2]}})
    with pytest.raises(ParserError, match="row and its header"):
        chunks(replace(parsed, blocks=[oversized]), original, small=True)


@pytest.mark.parametrize("kind", ["xlsx", "pptx"])
def test_actual_ingestion_preserves_office_locator_in_index_payload(tmp_path, kind):
    original = workbook_source(many=True) if kind == "xlsx" else presentation_source(long=True)
    path = tmp_path / original.filename
    path.write_bytes(original.content)
    with make_client(tmp_path) as client:
        response = client.post("/api/v1/ingestion/jobs", headers=web_transport_headers(
            actor_subject_id="user_dev", authorization_proof=True,
        ), json={"idempotency_key": f"office-source-locators-{kind}", "document_id": "doc_office",
                 "document_version_id": "ver_office", "source_file_uri": str(path), "ocr_enabled": False,
                 "expected_current_ingestion_job_id": None})
        assert response.status_code == 201, response.text
        report = client.get(f"/api/v1/ingestion/jobs/{response.json()['job_id']}/report", headers=actor_proof_headers()).json()
        assert report["status"] == "completed", report
        assert report["pages_processed"] == 0
        points = client.app.state.indexer.mock_points
        assert points
        for point in points:
            payload = point["payload"]
            assert payload["page_number"] is None
            assert payload["source_sha256"] == original.sha256
            locator = payload["metadata"]["source_locator"]
            assert locator["kind"] == ("sheet" if kind == "xlsx" else "slide")
            if payload["metadata"]["block_type"] == "table":
                assert "Řádky " in payload["section_path"][-1]
                assert len(locator["row_numbers"]) == len(payload["text"].splitlines())


def test_processing_limit_is_terminal_even_if_ocr_could_return_text(tmp_path, monkeypatch):
    import parsers.xlsx as xlsx_module
    monkeypatch.setattr(xlsx_module, "OfficeLimits", lambda: replace(OfficeLimits(), rows=10))
    settings = load_settings({"AKL_ENV": "test", "AKL_AUTH_MODE": "disabled", "AKL_INGESTION_DOCLING_MODE": "off"})
    router = ParserRouter(settings)
    def forbidden(*_args, **_kwargs):
        pytest.fail("A hard processing limit must not fall through to OCR")
    monkeypatch.setattr(router.ocr_provider, "extract", forbidden)
    with pytest.raises(ParserError) as error:
        router.parse(workbook_source(), parser_profile="default", ocr_enabled=True)
    assert error.value.code == "XLSX_PROCESSING_LIMIT"

    path = tmp_path / "budget.xlsx"
    path.write_bytes(workbook_source().content)
    with make_client(tmp_path) as client:
        response = client.post("/api/v1/ingestion/jobs", headers=web_transport_headers(
            actor_subject_id="user_dev", authorization_proof=True,
        ), json={"idempotency_key": "office-limit-no-partial-index", "document_id": "doc_limit",
                 "document_version_id": "ver_limit", "source_file_uri": str(path), "ocr_enabled": True,
                 "expected_current_ingestion_job_id": None})
        assert response.status_code == 201, response.text
        report = client.get(f"/api/v1/ingestion/jobs/{response.json()['job_id']}/report", headers=actor_proof_headers()).json()
        assert report["status"] == "failed"
        assert report["chunks_created"] == 0
        assert client.app.state.indexer.mock_points == []
        assert any(item["code"] == "XLSX_PROCESSING_LIMIT" for item in report["errors"])


def test_formula_and_merged_worksheet_content_remains_explicitly_review_required():
    from app.pipeline import _quality_report
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Description", "Calculated", "Value"])
    sheet.append(["Readable source evidence " * 30, "=1+2", 5])
    sheet.merge_cells("A4:C4")
    sheet.cell(4, 1, "Merged meaning")
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = XlsxParser().parse(source("formula.xlsx", buffer.getvalue()), parser_profile="default")
    assert {"XLSX_FORMULAS_REQUIRE_REVIEW", "XLSX_MERGED_CELLS_REQUIRE_REVIEW"} <= {code for code, _ in parsed.warnings}
    quality = _quality_report(parsed, extraction_profile="document_text_v1")
    assert quality.quality_tier == "good"  # Text density does not approve missing formula/merge semantics.
    assert quality.requires_review is True
    assert "=1+2" not in "\n".join(block.text for block in parsed.blocks)


def test_merged_presentation_cells_require_source_review():
    from app.pipeline import _quality_report
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(5), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged heading"
    table.cell(1, 2).text = "Readable source evidence " * 30
    buffer = io.BytesIO()
    presentation.save(buffer)
    parsed = PptxParser().parse(source("merged.pptx", buffer.getvalue()), parser_profile="default")
    assert "PPTX_MERGED_CELLS_REQUIRE_REVIEW" in {code for code, _ in parsed.warnings}
    assert _quality_report(parsed, extraction_profile="document_text_v1").requires_review is True
